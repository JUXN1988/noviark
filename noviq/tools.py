# -*- coding: utf-8 -*-
"""Herramientas que la IA puede usar (estilo Claude): archivos, proyectos, PowerShell, Python,
procesos en segundo plano, internet, imágenes, capturas, tareas y memoria."""
import base64
import difflib
import fnmatch
import html
import io
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
import webbrowser
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

import requests

from .config import (CONFIG, IS_WINDOWS, MEMORY_FILE, NO_WINDOW, OS_NAME, general_dir, meta_dir, meta_file, project_path, all_roots,
                     powershell_exe, projects_root)

MAX_OUT = 30000
SKIP_DIRS = {".git", "node_modules", ".venv", "venv", "__pycache__", ".idea", ".vs", "dist", "build", ".next",
             ".gradle", "bin", "obj", ".mypy_cache", ".pytest_cache"}
TEXT_EXT_BIN = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".ico", ".exe", ".dll", ".zip", ".7z", ".rar",
                ".mp3", ".mp4", ".wav", ".pdf", ".docx", ".xlsx", ".pptx", ".bin", ".so", ".pyc", ".class", ".jar"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


def cut(text, n=MAX_OUT):
    text = text or ""
    return text if len(text) <= n else text[:n // 2] + f"\n\n...[{len(text) - n} caracteres omitidos]...\n\n" + text[-n // 2:]


class ToolContext:
    def __init__(self, conv, emit, stop_ev, call_id=""):
        self.conv, self.emit, self.stop_ev, self.call_id = conv, emit, stop_ev, call_id

    @property
    def cwd(self) -> Path:
        pr = self.conv.get("project")
        if pr:
            p = project_path(pr)
            if p.exists():
                return p
        return general_dir()

    def resolve(self, ruta) -> Path:
        ruta = os.path.expandvars(os.path.expanduser(str(ruta or ".").strip().strip('"').strip("'")))
        p = Path(ruta)
        if not p.is_absolute():
            p = self.cwd / p
        return p.resolve()


def inside_root(p: Path) -> bool:
    for root in all_roots():
        try:
            p.resolve().relative_to(root.resolve())
            return True
        except (ValueError, OSError):
            continue
    return False


def file_event(ctx, p: Path, **extra):
    info = {"name": p.name, "path": str(p)}
    try:
        info["size"] = p.stat().st_size
    except OSError:
        pass
    info.update(extra)
    ctx.emit("file", info)


def img_to_dataurl(path: Path, max_side=1568):
    try:
        from PIL import Image
        im = Image.open(path)
        im.thumbnail((max_side, max_side))
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=85)
        return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
    except Exception:
        mime = {"png": "png", "jpg": "jpeg", "jpeg": "jpeg", "gif": "gif", "webp": "webp"}.get(path.suffix[1:].lower(), "png")
        return f"data:image/{mime};base64," + base64.b64encode(path.read_bytes()).decode()


# ═════════════════════════════ PROYECTOS ═════════════════════════════
TIPOS = {
    "python": ["src", "tests", "docs"],
    "web": ["css", "js", "img"],
    "node": ["src", "public", "tests"],
    "java": ["src/main/java", "src/test/java"],
    "csharp": ["src", "tests"],
    "arduino": ["docs"],
    "android": ["app/src/main/java", "app/src/main/res"],
    "datos": ["datos", "notebooks", "resultados"],
    "documentos": ["documentos", "recursos"],
    "general": [],
}
GITIGNORE = {
    "python": "__pycache__/\n*.pyc\n.venv/\nvenv/\n.env\ndist/\nbuild/\n*.egg-info/\n",
    "node": "node_modules/\ndist/\n.env\n*.log\n",
    "web": ".DS_Store\nThumbs.db\n",
    "java": "*.class\ntarget/\nbuild/\n.gradle/\n",
    "csharp": "bin/\nobj/\n.vs/\n",
    "android": ".gradle/\nbuild/\nlocal.properties\n*.apk\n",
    "datos": "__pycache__/\n.ipynb_checkpoints/\n.venv/\n",
}


def safe_name(n):
    n = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", (n or "").strip()).strip(". ")
    return n[:80] or "Proyecto"


def project_info(name):
    p = project_path(name)
    meta = {}
    f = meta_file(name)
    if f.exists():
        try:
            meta = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            pass
    ext = name in (CONFIG.get("proyectos_externos") or {})
    act = 0
    for cand in (p / ".noviq" / "ESTADO.md", p / ".noviq" / "bitacora.md", f, p):
        try:
            act = max(act, cand.stat().st_mtime)
        except OSError:
            pass
    return {"name": name, "path": str(p), "externo": ext, "existe": p.exists(), "actividad": act, **meta}


def list_projects():
    out, seen = [], set()
    root = projects_root()
    for d in root.iterdir():
        if d.is_dir() and not d.name.startswith((".", "_")):
            out.append(project_info(d.name))
            seen.add(d.name)
    for name in (CONFIG.get("proyectos_externos") or {}):
        if name not in seen:
            out.append(project_info(name))
    return sorted(out, key=lambda x: -x.get("actividad", 0))


def detect_project(path: Path):
    """Detecta el tipo de un proyecto existente y cómo ejecutarlo/probarlo."""
    has = lambda *g: any(next(path.glob(x), None) for x in g)
    tipo, run, test = "general", "", ""
    if has("pubspec.yaml"):
        tipo, run, test = "flutter", "flutter run", "flutter test"
    elif has("AndroidManifest.xml", "app/src/main/AndroidManifest.xml") or (has("build.gradle*") and has("app")):
        tipo, run, test = "android", ".\\gradlew assembleDebug", ".\\gradlew test"
    elif has("package.json"):
        tipo, run, test = "node", "npm start", "npm test"
        try:
            pk = json.loads((path / "package.json").read_text(encoding="utf-8"))
            sc = pk.get("scripts", {})
            run = "npm run dev" if "dev" in sc else ("npm start" if "start" in sc else run)
            test = "npm test" if "test" in sc else ""
        except Exception:
            pass
    elif has("*.sln", "*.csproj"):
        tipo, run, test = "csharp", "dotnet run", "dotnet test"
    elif has("pom.xml"):
        tipo, run, test = "java", "mvn exec:java", "mvn test"
    elif has("build.gradle*"):
        tipo, run, test = "java", "gradle run", "gradle test"
    elif has("*.ino", "*/*.ino", "platformio.ini"):
        tipo = "arduino"
    elif has("pyproject.toml", "requirements.txt", "setup.py", "*.py", "src/*.py"):
        tipo = "python"
        main = next((n for n in ("main.py", "app.py", "src/main.py", "run.py", "manage.py") if (path / n).exists()), "")
        run = f"python {main}" if main else ""
        test = "python -m pytest -q" if has("tests", "test_*.py", "tests/*.py") else ""
    elif has("index.html", "*.html"):
        tipo, run = "web", "start index.html"
    elif has("composer.json", "*.php"):
        tipo, run = "php", "php -S localhost:8000"
    return {"tipo": tipo, "ejecutar": run, "probar": test}


def import_project(path_str, nombre=""):
    p = Path(os.path.expandvars(os.path.expanduser(path_str.strip().strip('"')))).resolve()
    if not p.is_dir():
        raise ValueError(f"No existe la carpeta: {p}")
    name = safe_name(nombre or p.name)
    ext = dict(CONFIG.get("proyectos_externos") or {})
    base, n = name, 2
    while (name in ext and Path(ext[name]) != p) or (projects_root() / name).exists() and (projects_root() / name).resolve() != p:
        name, n = f"{base}_{n}", n + 1
    if (projects_root() / name).resolve() != p:
        ext[name] = str(p)
        CONFIG["proyectos_externos"] = ext
        from .config import save_config
        save_config()
    det = detect_project(p)
    f = meta_file(name)
    meta = {}
    if f.exists():
        try:
            meta = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            meta = {}
    meta.setdefault("tipo", det["tipo"])
    meta.setdefault("descripcion", "Proyecto existente importado de " + str(p))
    meta["comandos"] = {"ejecutar": det["ejecutar"], "probar": det["probar"]}
    meta.setdefault("importado", datetime.now().isoformat(timespec="seconds"))
    try:
        f.write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    except OSError:
        pass
    return name, p, det


def unlink_project(name):
    ext = dict(CONFIG.get("proyectos_externos") or {})
    if name in ext:
        ext.pop(name)
        CONFIG["proyectos_externos"] = ext
        from .config import save_config
        save_config()
        return True
    return False


def ps_quote(s):
    return "'" + str(s).replace("'", "''") + "'"


def create_project(nombre, tipo="general", descripcion="", instrucciones=""):
    nombre = safe_name(nombre)
    tipo = tipo if tipo in TIPOS else "general"
    root = projects_root()
    p = root / nombre
    subdirs = TIPOS[tipo]
    ps = powershell_exe()
    log = ""
    if ps:  # crea las carpetas con PowerShell
        script = (f"$p = Join-Path {ps_quote(root)} {ps_quote(nombre)}; "
                  f"New-Item -ItemType Directory -Force -Path $p | Out-Null; "
                  + "".join(f"New-Item -ItemType Directory -Force -Path (Join-Path $p {ps_quote(d)}) | Out-Null; " for d in subdirs)
                  + "Write-Output \"Carpeta creada: $p\"")
        r = subprocess.run([ps, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-Command", script],
                           capture_output=True, text=True, timeout=60, creationflags=NO_WINDOW)
        log = (r.stdout + r.stderr).strip()
    p.mkdir(parents=True, exist_ok=True)
    for d in subdirs:
        (p / d).mkdir(parents=True, exist_ok=True)
    meta = {"tipo": tipo, "descripcion": descripcion, "instrucciones": instrucciones,
            "creado": datetime.now().isoformat(timespec="seconds")}
    meta_file(nombre).write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    if not (p / "README.md").exists():
        (p / "README.md").write_text(f"# {nombre}\n\n{descripcion}\n", encoding="utf-8")
    if tipo in GITIGNORE and not (p / ".gitignore").exists():
        (p / ".gitignore").write_text(GITIGNORE[tipo], encoding="utf-8")
    git = shutil.which("git")
    if git and tipo not in ("documentos", "general") and not (p / ".git").exists():
        subprocess.run([git, "init", "-q"], cwd=str(p), capture_output=True, creationflags=NO_WINDOW)
        log += "\nRepositorio git inicializado."
    return p, log


def t_crear_proyecto(a, ctx):
    p, log = create_project(a.get("nombre"), a.get("tipo", "general"), a.get("descripcion", ""), a.get("instrucciones", ""))
    ctx.conv["project"] = p.name
    ctx.emit("project", {"name": p.name, "path": str(p)})
    tree = "\n".join("  " + str(x.relative_to(p)) + ("/" if x.is_dir() else "") for x in sorted(p.rglob("*"))
                     if ".git" not in x.parts)
    return f"Proyecto creado en {p}\n{log}\nEstructura:\n{tree}\n\nEste es ahora el directorio de trabajo: usa rutas relativas."


def t_cambiar_proyecto(a, ctx):
    n = a.get("nombre", "")
    names = [x["name"] for x in list_projects()]
    match = next((x for x in names if x.lower() == n.lower()), None) or next((x for x in names if n.lower() in x.lower()), None)
    if not match:
        return "No existe ese proyecto. Proyectos disponibles:\n" + ("\n".join(names) or "(ninguno)")
    ctx.conv["project"] = match
    ctx.emit("project", {"name": match, "path": str(project_path(match))})
    return f"Proyecto activo: {match} ({project_path(match)})\n\n" + t_listar_directorio({"ruta": ".", "recursivo": True}, ctx)


# ═════════════════════════════ ARCHIVOS ═════════════════════════════
def _backup(ctx, p: Path, data=None):
    """Guarda una copia antes de modificar un archivo (para deshacer_cambio). data = contenido anterior ya leído."""
    if data is None and (not p.exists() or not p.is_file() or p.stat().st_size > 5_000_000):
        return None
    base = meta_dir(ctx.conv["project"]) if ctx.conv.get("project") else general_dir() / ".noviq"
    folder = base / "respaldos"
    folder.mkdir(parents=True, exist_ok=True)
    hist = ctx.conv.setdefault("respaldos", [])
    prev = next((h for h in reversed(hist) if h.get("ruta") == str(p)), None)
    try:  # si la última copia es idéntica no se duplica (así deshacer_cambio vuelve a la versión anterior real)
        cur = data if data is not None else p.read_bytes()
        if prev and Path(prev["copia"]).exists() and Path(prev["copia"]).read_bytes() == cur:
            return Path(prev["copia"])
    except OSError:
        cur = None
    dest = folder / f"{p.name}.{datetime.now():%Y%m%d_%H%M%S_%f}.bak"
    if data is not None:
        dest.write_bytes(data)
    else:
        shutil.copy2(p, dest)
    hist.append({"ruta": str(p), "copia": str(dest), "t": datetime.now().isoformat(timespec="seconds")})
    del hist[:-200]
    return dest


PLACEHOLDER_RX = re.compile(r"(?im)^\s*(?://|#|/\*|<!--)\s*(?:\.\.\.|…|(?:el\s+)?resto\s+del\s+c[oó]digo|c[oó]digo\s+(?:anterior|existente|igual)|"
                            r"mismo\s+c[oó]digo|sin\s+cambios|rest\s+of\s+(?:the\s+)?code|existing\s+code|same\s+as\s+before)")


def _fix_escapes(contenido):
    """La IA mandó el archivo con \\n literales en vez de saltos de línea reales (JSON escapado dos veces)."""
    if isinstance(contenido, str) and "\n" not in contenido and contenido.count("\\n") >= 3:
        return contenido.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\\t", "\t").replace('\\"', '"'), True
    return contenido, False


def t_escribir_archivo(a, ctx):
    p = ctx.resolve(a.get("ruta", "archivo.txt"))
    p.parent.mkdir(parents=True, exist_ok=True)
    contenido = a.get("contenido", "")
    if not isinstance(contenido, str):
        contenido = json.dumps(contenido, ensure_ascii=False, indent=2) if isinstance(contenido, (dict, list)) else str(contenido)
    contenido, arreglado = _fix_escapes(contenido)
    notas = ["convertí los \\n literales en saltos de línea reales"] if arreglado else []
    existed = p.exists()
    crlf = p.suffix.lower() in (".bat", ".cmd", ".ps1")
    if a.get("modo") == "agregar" and existed:
        _backup(ctx, p)
        prev = _read_text(p)
        body = (prev if prev.endswith("\n") or not prev else prev + "\n") + contenido
        p.write_text(body, encoding="utf-8", newline="\r\n" if crlf else "\n")
        lines = body.count("\n") + 1
        file_event(ctx, p, lines=lines, action="ampliado")
        return f"Contenido agregado al final de {p} (ahora {lines} líneas)" + (f" ({'; '.join(notas)})" if notas else "")
    old_lines = 0
    if existed:
        old_lines = _read_text(p).count("\n") + 1
        if not a.get("sobrescribir") and old_lines > 60:
            return (f"ERROR: {p.name} ya existe y tiene {old_lines} líneas. Para CORREGIR o cambiar partes usa "
                    "editar_archivo (solo en las líneas afectadas; puedes enviar varias ediciones en una llamada). "
                    "Si de verdad necesitas reemplazarlo completo, repite la llamada con sobrescribir=true.")
        _backup(ctx, p)
    p.write_text(contenido, encoding="utf-8", newline="\r\n" if crlf else "\n")
    lines = contenido.count("\n") + 1
    file_event(ctx, p, lines=lines, action="reemplazado" if existed else "creado")
    avisos = []
    if existed and old_lines >= 40 and lines < old_lines * 0.4:
        avisos.append(f"⚠ El archivo pasó de {old_lines} a {lines} líneas: ¿borraste código sin querer? Si fue un error usa "
                      "deshacer_cambio (hay copia de respaldo).")
    m = PLACEHOLDER_RX.search(contenido)
    if m and p.suffix.lower() not in (".md", ".txt"):
        ln = contenido.count("\n", 0, m.start()) + 1
        avisos.append(f"⚠ La línea {ln} («{m.group(0).strip()[:50]}») es un marcador de código OMITIDO: el archivo quedó incompleto. "
                      "Escribe ese código de verdad (con editar_archivo en esa línea).")
    return (f"Archivo {'reemplazado (copia de respaldo guardada)' if existed else 'creado'}: {p} ({lines} líneas)"
            + (f" ({'; '.join(notas)})" if notas else "") + ("\n" + "\n".join(avisos) if avisos else ""))


def _read_text(p: Path):
    from . import texto as TX
    return TX.read_text(p)


def t_editar_archivo(a, ctx):
    from . import edicion as ED
    from . import texto as TX
    p = ctx.resolve(a.get("ruta", ""))
    if not p.exists():
        return f"ERROR: no existe {p}. Usa escribir_archivo para crearlo."
    if p.is_dir():
        return f"ERROR: {p} es una carpeta, no un archivo."
    orig, enc = TX.read(p)
    crlf = "\r\n" in orig
    text = orig.replace("\r\n", "\n")
    result, total, notas = text, 0, []
    ediciones = a.get("ediciones") or []
    if isinstance(ediciones, str):
        try:
            ediciones = json.loads(ediciones)
        except Exception:
            ediciones = []
    if isinstance(ediciones, dict):
        ediciones = [ediciones]
    rango = None
    if a.get("desde_linea"):
        try:
            rango = (int(a["desde_linea"]), int(a.get("hasta_linea") or a["desde_linea"]))
        except (TypeError, ValueError):
            rango = None
    if a.get("texto_anterior"):
        ediciones = [{"texto_anterior": a.get("texto_anterior"), "texto_nuevo": a.get("texto_nuevo", ""),
                      "reemplazar_todo": a.get("reemplazar_todo")}] + list(ediciones)
    if ediciones:
        for i, ed in enumerate(ediciones, 1):
            if not isinstance(ed, dict):
                return f"ERROR en la edición {i}: cada edición debe ser un objeto {{texto_anterior, texto_nuevo}}."
            result, n, err, nota = ED.aplicar(result, ed.get("texto_anterior"), ed.get("texto_nuevo"),
                                               bool(ed.get("reemplazar_todo")), rango if len(ediciones) == 1 else None)
            if err:
                return f"ERROR en la edición {i} de {len(ediciones)} (no se modificó nada): {err}"
            if nota:
                notas.append(nota)
            total += n
    elif a.get("despues_de_linea") is not None and str(a.get("despues_de_linea")).strip() != "":
        lines = result.split("\n")
        try:
            k = int(a["despues_de_linea"])
        except (TypeError, ValueError):
            return "ERROR: despues_de_linea debe ser un número de línea (0 = al principio)."
        if k < 0 or k > len(lines):
            return f"ERROR: el archivo tiene {len(lines)} líneas; despues_de_linea debe estar entre 0 y {len(lines)}."
        nuevo = (a.get("texto_nuevo") or "").replace("\r\n", "\n")
        lines[k:k] = nuevo.split("\n")
        result, total = "\n".join(lines), 1
    elif rango:
        lines = result.split("\n")
        d, h = max(1, rango[0]), min(len(lines), max(rango))
        if d > len(lines):
            return f"ERROR: el archivo tiene {len(lines)} líneas (pediste desde la {d})."
        nuevo = (a.get("texto_nuevo") or "").replace("\r\n", "\n")
        lines[d - 1:h] = nuevo.split("\n") if nuevo != "" else []
        result, total = "\n".join(lines), 1
    else:
        return ("ERROR: indica texto_anterior + texto_nuevo, una lista 'ediciones', desde_linea/hasta_linea + texto_nuevo "
                "(reemplazar ese rango) o despues_de_linea + texto_nuevo (insertar).")
    if result == text:
        return "Sin cambios: el texto nuevo es igual al anterior."
    _backup(ctx, p)
    p.write_text(result.replace("\n", "\r\n") if crlf else result, encoding="utf-8", newline="")
    if not TX.es_utf8(enc):
        notas.append(f"el archivo estaba en {TX.nombre(enc)}; ahora quedó en UTF-8")
    diff = "".join(difflib.unified_diff(text.splitlines(True), result.splitlines(True), "antes", "después", n=1))
    file_event(ctx, p, action="editado", diff=cut(diff, 6000))
    return (f"Editado {p} ({total} cambio(s); copia de respaldo guardada)." + (" Nota: " + "; ".join(notas) + "." if notas else "")
            + f"\n{cut(diff, 3000)}")


def t_deshacer_cambio(a, ctx):
    p = ctx.resolve(a.get("ruta", ""))
    hist = [h for h in ctx.conv.get("respaldos", []) if Path(h["ruta"]) == p and Path(h["copia"]).exists()]
    if hist:
        last = hist[-1]
        shutil.copy2(last["copia"], p)
        ctx.conv["respaldos"].remove(last)
        Path(last["copia"]).unlink(missing_ok=True)
        file_event(ctx, p, action="restaurado")
        return f"Restaurado {p} a la versión del {last['t']}. Quedan {len(hist) - 1} copias anteriores."
    # copias de otras conversaciones (carpeta .noviq/respaldos del proyecto)
    base = meta_dir(ctx.conv["project"], create=False) if ctx.conv.get("project") else general_dir() / ".noviq"
    folder = base / "respaldos"
    cands = sorted(folder.glob(p.name + ".*.bak"), reverse=True) if folder.exists() else []
    if cands:
        dest = cands[0]
        cur = _backup(ctx, p) if p.exists() else None
        shutil.copy2(dest, p)
        file_event(ctx, p, action="restaurado")
        return (f"Restaurado {p} con la copia {dest.name} (de un trabajo anterior)."
                + (f" La versión que había quedó guardada como {cur.name}." if cur else ""))
    return (f"ERROR: no hay copias de respaldo de {p}. Prueba puntos_restauracion (accion='listar' y luego 'volver' con "
            "archivos=[...]) para recuperar una versión anterior, o reescríbelo completo con escribir_archivo.")


# ───────── Localizar errores y mapa del código ─────────
ERR_PATTERNS = [
    re.compile(r'File "([^"]+)", line (\d+)'),                                    # Python
    re.compile(r'at (?:.*? \()?((?:[A-Za-z]:)?[^\s():]+\.\w+):(\d+)(?::\d+)?\)?'),     # Node/JS
    re.compile(r'((?:[A-Za-z]:)?[\w./\\ -]+\.(?:ts|tsx|cs|vb|cpp|c|h))\((\d+),\d+\)'),    # TS / MSBuild
    re.compile(r'\((\w+\.(?:java|kt)):(\d+)\)'),                                      # Java / Kotlin
    re.compile(r'At ((?:[A-Za-z]:)?[^\n:]+\.ps1):(\d+) char'),                         # PowerShell
    re.compile(r' in ((?:[A-Za-z]:)?[^\s]+\.php) on line (\d+)'),                      # PHP
    re.compile(r' in ((?:[A-Za-z]:)?[^\s:]+\.\w+):line (\d+)'),                          # .NET
    re.compile(r'((?:[A-Za-z]:)?[\w./\\-]+\.(?:py|js|jsx|ts|tsx|html|css|java|cs|go|rs|rb|php|ino|cpp|c|kt|dart|lua)):(\d+)'),
]


def _find_file(ctx, name):
    p = ctx.resolve(name)
    if p.exists():
        return p
    base = Path(name.replace("\\", "/")).name
    for f in _iter_files(ctx.cwd, base):
        return f
    return None


def t_localizar_error(a, ctx):
    err = a.get("texto_error") or ""
    found, seen = [], set()
    for rx in ERR_PATTERNS:
        for m in rx.finditer(err):
            f, ln = m.group(1).strip(), int(m.group(2))
            if "site-packages" in f or "node_modules" in f or f.startswith("<") or "lib/python" in f.lower():
                continue
            p = _find_file(ctx, f)
            if p and (str(p), ln) not in seen:
                seen.add((str(p), ln))
                found.append((p, ln))
    out = []
    for p, ln in found[-4:]:  # los últimos marcos suelen ser el código del usuario
        lines = _read_text(p).split("\n")
        a0, b0 = max(1, ln - 12), min(len(lines), ln + 12)
        snippet = "\n".join(f"{'>>' if i == ln else '  '}{i:>5}\t{lines[i - 1]}" for i in range(a0, b0 + 1))
        out.append(f"### {p} (línea {ln})\n{snippet}")
    if not out:
        # sin ruta en el error: busca identificadores citados en el mensaje
        ids = set(re.findall(r"['\"`]([A-Za-z_][\w.]{2,60})['\"`]", err))
        for ident in list(ids)[:5]:
            r = t_buscar_en_archivos({"regex": re.escape(ident.split(".")[-1]), "ruta": "."}, ctx)
            if not r.startswith("Sin coincidencias"):
                out.append(f"### Coincidencias de «{ident}»\n" + "\n".join(r.splitlines()[:15]))
    if not out:
        return ("No encontré archivos ni líneas del proyecto en ese error. Usa mapa_codigo y buscar_en_archivos "
                "para ubicar la parte relacionada.")
    return ("Ubicación probable del error (>> marca la línea). Corrige SOLO estas partes con editar_archivo:\n\n"
            + "\n\n".join(out))


DEF_RX = {
    ".py": re.compile(r"^\s*(?:async\s+)?(def|class)\s+(\w+)"),
    ".js": re.compile(r"^\s*(?:export\s+)?(?:default\s+)?(?:async\s+)?(function|class)\s+(\w+)|^\s*(?:export\s+)?(?:const|let|var)\s+(\w+)\s*=\s*(?:async\s*)?(?:\([^)]*\)|\w+)\s*=>"),
    ".cs": re.compile(r"^\s*(?:public|private|protected|internal|static|\s)*\s*(class|interface|enum|void|[\w<>\[\]]+)\s+(\w+)\s*[({]"),
    ".java": re.compile(r"^\s*(?:public|private|protected|static|final|\s)*\s*(class|interface|enum|void|[\w<>\[\]]+)\s+(\w+)\s*[({]"),
    ".php": re.compile(r"^\s*(?:public|private|protected|static|\s)*(function|class)\s+(\w+)"),
    ".ps1": re.compile(r"^\s*(function)\s+([\w-]+)"),
    ".go": re.compile(r"^\s*(func|type)\s+(?:\([^)]*\)\s*)?(\w+)"),
    ".rs": re.compile(r"^\s*(?:pub\s+)?(fn|struct|enum|impl|trait)\s+(\w+)"),
    ".kt": re.compile(r"^\s*(?:\w+\s+)*(fun|class|object|interface)\s+(\w+)"),
    ".dart": re.compile(r"^\s*(class)\s+(\w+)|^\s*[\w<>?]+\s+(\w+)\s*\([^;]*\)\s*(?:async\s*)?\{"),
    ".ino": re.compile(r"^\s*(void|int|float|bool|String|class)\s+(\w+)\s*\("),
}
DEF_RX.update({".ts": DEF_RX[".js"], ".tsx": DEF_RX[".js"], ".jsx": DEF_RX[".js"], ".mjs": DEF_RX[".js"],
               ".cpp": DEF_RX[".ino"], ".c": DEF_RX[".ino"], ".h": DEF_RX[".ino"]})


def t_mapa_codigo(a, ctx):
    root = ctx.resolve(a.get("ruta") or ".")
    files = [root] if root.is_file() else list(_iter_files(root))
    out, n = [], 0
    for f in sorted(files):
        rx = DEF_RX.get(f.suffix.lower())
        if not rx:
            if f.suffix.lower() in (".html", ".css", ".json", ".md", ".txt", ".bat", ".toml", ".yml", ".yaml", ".ini", ".xml", ".sql"):
                try:
                    out.append(f"{f.relative_to(root) if root.is_dir() else f.name} ({_read_text(f).count(chr(10)) + 1} líneas)")
                except (OSError, ValueError):
                    pass
            continue
        try:
            lines = _read_text(f).split("\n")
        except OSError:
            continue
        defs = []
        for i, line in enumerate(lines, 1):
            m = rx.match(line)
            if m:
                name = next((g for g in m.groups()[1:] if g), None) or m.group(1)
                if name and name not in ("if", "for", "while", "switch", "return", "catch", "new"):
                    defs.append(f"{name}@{i}")
        rel = f.relative_to(root) if root.is_dir() else f.name
        out.append(f"{rel} ({len(lines)} líneas): " + (", ".join(defs[:60]) or "—"))
        n += 1
        if n >= 200:
            out.append("... (más archivos)")
            break
    return cut("Mapa del código (nombre@línea):\n" + "\n".join(out), 20000) if out else "No hay archivos de código."


def t_leer_archivo(a, ctx):
    p = ctx.resolve(a.get("ruta", ""))
    if not p.exists():
        return f"ERROR: no existe {p}"
    if p.is_dir():
        return t_listar_directorio({"ruta": str(p)}, ctx)
    ext = p.suffix.lower()
    if ext in IMG_EXT:
        ctx.emit("image", {"name": p.name, "path": str(p)})
        return {"text": f"Imagen {p.name} adjuntada para que la veas.", "images": [img_to_dataurl(p)]}
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            rd = PdfReader(str(p))
            parts = [f"--- Página {i + 1} ---\n{(pg.extract_text() or '').strip()}" for i, pg in enumerate(rd.pages[:80])]
            return cut(f"PDF {p.name}: {len(rd.pages)} páginas\n\n" + "\n\n".join(parts), 80000)
        except ImportError:
            return "ERROR: falta el paquete pypdf (pip install pypdf)."
    if ext == ".docx":
        try:
            import docx
            d = docx.Document(str(p))
            txt = "\n".join(par.text for par in d.paragraphs)
            for ti, t in enumerate(d.tables):
                txt += f"\n\n[Tabla {ti + 1}]\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in t.rows)
            return cut(txt, 80000)
        except ImportError:
            return "ERROR: falta el paquete python-docx."
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"=== Hoja: {ws.title} ===")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 300:
                        out.append("... (más filas)")
                        break
                    out.append(" | ".join("" if v is None else str(v) for v in row))
            return cut("\n".join(out), 80000)
        except ImportError:
            return "ERROR: falta el paquete openpyxl."
    if ext in TEXT_EXT_BIN:
        return f"ERROR: {p.name} es un archivo binario ({p.stat().st_size} bytes)."
    from . import texto as TX
    data = p.read_bytes()
    text, enc = TX.decode(data)
    lines = text.splitlines()
    try:
        a0 = max(1, int(a.get("desde_linea") or 1))
        b0 = min(len(lines), int(a.get("hasta_linea") or (a0 + 1999)))
    except (TypeError, ValueError):
        a0, b0 = 1, min(len(lines), 2000)
    if lines and a0 > len(lines):
        return f"ERROR: {p.name} solo tiene {len(lines)} líneas (pediste desde la {a0}). Lee con desde_linea entre 1 y {len(lines)}."
    body = "\n".join(f"{i:>5}\t{lines[i - 1]}" for i in range(a0, b0 + 1))
    more = f"\n... ({len(lines) - b0} líneas más; usa desde_linea={b0 + 1})" if b0 < len(lines) else ""
    head = f"{p} ({len(lines)} líneas)"
    probs = TX.problemas(data, text, enc)
    if probs:
        head += "\n⚠ " + "; ".join(probs) + ". (editar_archivo/escribir_archivo lo guardan en UTF-8; diagnosticar lo repara solo.)"
    if not lines:
        head += "\n(el archivo está VACÍO)"
    return cut(f"{head}\n{body}{more}", 100000)


def t_listar_directorio(a, ctx):
    p = ctx.resolve(a.get("ruta") or ".")
    if not p.exists():
        return f"ERROR: no existe {p}"
    rec = bool(a.get("recursivo"))
    out, n = [], 0

    def walk(d, depth):
        nonlocal n
        try:
            items = sorted(d.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))
        except OSError as e:
            out.append("  " * depth + f"[sin acceso: {e}]")
            return
        for c in items:
            if n >= 400:
                return
            n += 1
            if c.is_dir():
                out.append("  " * depth + f"📁 {c.name}/")
                if rec and depth < 4 and c.name not in SKIP_DIRS:
                    walk(c, depth + 1)
            else:
                try:
                    sz = c.stat().st_size
                except OSError:
                    sz = 0
                out.append("  " * depth + f"   {c.name}  ({sz:,} B)")

    walk(p, 0)
    if n >= 400:
        out.append("... (lista recortada)")
    return f"{p}\n" + ("\n".join(out) or "(vacío)")


def _iter_files(root: Path, pattern="*"):
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if d not in SKIP_DIRS and not d.startswith(".")]
        for f in fns:
            rel = os.path.relpath(os.path.join(dp, f), root)
            if fnmatch.fnmatch(f, pattern) or fnmatch.fnmatch(rel.replace("\\", "/"), pattern):
                yield Path(dp) / f


def t_buscar_archivos(a, ctx):
    root = ctx.resolve(a.get("ruta") or ".")
    pat = a.get("patron") or "*"
    pat = pat[3:] if pat.startswith("**/") else pat
    res = []
    for f in _iter_files(root, pat):
        res.append(str(f.relative_to(root)) if str(f).startswith(str(root)) else str(f))
        if len(res) >= 300:
            res.append("... (más resultados)")
            break
    return f"{len(res)} archivo(s) en {root}:\n" + "\n".join(res) if res else f"Sin coincidencias para '{pat}' en {root}"


def t_buscar_en_archivos(a, ctx):
    root = ctx.resolve(a.get("ruta") or ".")
    try:
        rx = re.compile(a.get("regex", ""), 0 if a.get("mayusculas") else re.I)
    except re.error as e:
        return f"ERROR en la expresión regular: {e}"
    files = [root] if root.is_file() else _iter_files(root, a.get("incluir") or "*")
    out = []
    for f in files:
        if f.suffix.lower() in TEXT_EXT_BIN:
            continue
        try:
            if f.stat().st_size > 3_000_000:
                continue
            lines = _read_text(f).splitlines()
            ctxn = min(int(a.get("contexto") or 0), 6)
            for i, line in enumerate(lines, 1):
                if rx.search(line):
                    rel = f.relative_to(root) if root.is_dir() else f.name
                    if ctxn:
                        blk = [f"{rel}:{j}{'>' if j == i else '-'} {lines[j - 1].rstrip()[:200]}"
                               for j in range(max(1, i - ctxn), min(len(lines), i + ctxn) + 1)]
                        out.append("\n".join(blk) + "\n--")
                    else:
                        out.append(f"{rel}:{i}: {line.strip()[:240]}")
                    if len(out) >= 250:
                        break
        except (OSError, ValueError):
            continue
        if len(out) >= 250:
            out.append("... (más coincidencias)")
            break
    return "\n".join(out) if out else "Sin coincidencias."


def t_mover_o_copiar(a, ctx):
    src, dst = ctx.resolve(a.get("origen", "")), ctx.resolve(a.get("destino", ""))
    if not src.exists():
        return f"ERROR: no existe {src}"
    if dst.is_dir():
        dst = dst / src.name
    dst.parent.mkdir(parents=True, exist_ok=True)
    if a.get("copiar"):
        (shutil.copytree if src.is_dir() else shutil.copy2)(src, dst)
        return f"Copiado: {src} → {dst}"
    shutil.move(str(src), str(dst))
    return f"Movido: {src} → {dst}"


def t_eliminar(a, ctx):
    p = ctx.resolve(a.get("ruta", ""))
    if not p.exists():
        return f"ERROR: no existe {p}"
    if IS_WINDOWS:
        kind = "DeleteDirectory" if p.is_dir() else "DeleteFile"
        script = (f"Add-Type -AssemblyName Microsoft.VisualBasic; [Microsoft.VisualBasic.FileIO.FileSystem]::{kind}("
                  f"{ps_quote(p)}, 'OnlyErrorDialogs', 'SendToRecycleBin')")
        r = subprocess.run(["powershell", "-NoProfile", "-Command", script], capture_output=True, text=True,
                           timeout=60, creationflags=NO_WINDOW)
        if p.exists():
            return f"ERROR al enviar a la papelera: {r.stderr[:500]}"
        return f"Enviado a la Papelera de reciclaje: {p}"
    (shutil.rmtree if p.is_dir() else os.remove)(p)
    return f"Eliminado: {p}"


# ═════════════════════════════ COMANDOS ═════════════════════════════
PROCS = {}  # id -> {"proc", "log", "cmd", "cwd", "start"}


def _kill_tree(proc):
    try:
        if IS_WINDOWS:
            subprocess.run(["taskkill", "/T", "/F", "/PID", str(proc.pid)], capture_output=True, creationflags=NO_WINDOW)
        else:
            proc.kill()
    except Exception:
        pass


def _env():
    return dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1", PYTHONUNBUFFERED="1")


def run_streaming(cmd, ctx, timeout=300):
    """Ejecuta mostrando la salida en vivo en el chat."""
    proc = subprocess.Popen(cmd, cwd=str(ctx.cwd), stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                            stdin=subprocess.DEVNULL, env=_env(), creationflags=NO_WINDOW)
    chunks, pending, last = [], [], [time.time()]

    def reader():
        for raw in iter(proc.stdout.readline, b""):
            line = raw.decode("utf-8", errors="replace")
            chunks.append(line)
            pending.append(line)
            if time.time() - last[0] > 0.3:
                ctx.emit("tool_output", {"id": ctx.call_id, "text": "".join(pending)})
                pending.clear()
                last[0] = time.time()

    th = threading.Thread(target=reader, daemon=True)
    th.start()
    t0, status = time.time(), None
    while proc.poll() is None:
        if ctx.stop_ev.is_set():
            _kill_tree(proc)
            status = "DETENIDO por el usuario"
            break
        if time.time() - t0 > timeout:
            _kill_tree(proc)
            status = f"TIEMPO AGOTADO ({timeout}s). Si es un servidor o proceso largo usa segundo_plano=true."
            break
        time.sleep(0.1)
    th.join(timeout=3)
    if pending:
        ctx.emit("tool_output", {"id": ctx.call_id, "text": "".join(pending)})
    out = "".join(chunks).strip()
    code = proc.returncode
    head = status or f"Código de salida: {code}"
    return cut(f"{head}\n{out if out else '(sin salida)'}")


def run_background(cmd, ctx, label):
    pid = uuid.uuid4().hex[:6]
    logdir = ctx.cwd / ".noviq_logs"
    logdir.mkdir(exist_ok=True)
    log = logdir / f"proceso_{pid}.log"
    fh = open(log, "wb")
    proc = subprocess.Popen(cmd, cwd=str(ctx.cwd), stdout=fh, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL,
                            env=_env(), creationflags=NO_WINDOW)
    PROCS[pid] = {"proc": proc, "log": log, "cmd": label, "cwd": str(ctx.cwd), "start": time.time(), "fh": fh}
    time.sleep(2.5)
    first = log.read_bytes().decode("utf-8", errors="replace")[-3000:]
    state = "en ejecución" if proc.poll() is None else f"terminó con código {proc.returncode}"
    return f"Proceso en segundo plano id={pid} ({state}).\nPrimeras líneas:\n{first or '(sin salida aún)'}\n\nUsa procesos(accion='salida', id='{pid}') para ver más."


def ps_command(comando):
    ps = powershell_exe()
    if ps:
        prefix = ("[Console]::OutputEncoding=[System.Text.Encoding]::UTF8; $OutputEncoding=[System.Text.Encoding]::UTF8; "
                  "$PSDefaultParameterValues['*:Encoding']='utf8'; $ProgressPreference='SilentlyContinue'; ")
        return [ps, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-Command", prefix + comando]
    return ["bash", "-lc", comando]


CODE_FILE_RX = r"[\w./\\:-]+\.(?:html?|js|mjs|jsx|ts|tsx|css|py|json|java|cs|php|vue|kt|dart|cpp|c|h|go|rs|rb|lua|sql|xml|md|txt|ps1|bat|ini|ya?ml|toml)\b"
WRITE_CMD_RX = re.compile(r"(?i)\b(set-content|add-content|out-file|new-item\b[^\n|;]*-value)\b|\[(?:system\.)?io\.file\]::write(?:all)?(?:text|lines|bytes)|"
                          r"(?<![<>=!-])>{1,2}\s*['\"]?" + CODE_FILE_RX)


def shell_file_write(comando):
    """¿El comando escribe el CONTENIDO de un archivo de código desde el propio comando (here-string o texto largo)?"""
    c = comando or ""
    if not WRITE_CMD_RX.search(c) or not re.search(CODE_FILE_RX, c, re.I):
        return False
    heredoc = re.search(r"@['\"]\s*\r?\n", c) or re.search(r"<<\s*['\"]?\w+", c)
    return bool(heredoc) or len(c) > 900


def _mentioned_files(ctx, texto):
    """Archivos existentes del proyecto que se nombran en un comando/código (para respaldarlos antes)."""
    out, seen = [], set()
    for m in re.finditer(CODE_FILE_RX, texto or "", re.I):
        tok = m.group(0).strip("'\"").rstrip(".,;:)")
        try:
            pth = ctx.resolve(tok)
        except (OSError, ValueError):
            continue
        if pth in seen or not pth.is_file() or not inside_root(pth):
            continue
        seen.add(pth)
        out.append(pth)
        if len(out) >= 12:
            break
    return out


def _pre_backup(ctx, texto):
    """Lee (en memoria) los archivos del proyecto que nombra un comando, para respaldar después SOLO los que cambien."""
    pre = {}
    for pth in _mentioned_files(ctx, texto):
        try:
            if pth.stat().st_size <= 3_000_000:
                pre[pth] = pth.read_bytes()
        except OSError:
            pass
    return pre


def _post_backup(ctx, pre):
    cambiados = []
    for pth, data in pre.items():
        try:
            if not pth.exists() or pth.read_bytes() != data:
                _backup(ctx, pth, data)
                cambiados.append(pth.name)
        except OSError:
            pass
    return cambiados


def t_ejecutar_powershell(a, ctx):
    comando = a.get("comando", "")
    if shell_file_write(comando) and not a.get("forzar"):
        return ("ERROR (bloqueado para proteger tus archivos): no escribas archivos de código con PowerShell "
                "(Set-Content/Out-File/here-strings @\"…\"@). Rompe las comillas y las $variables del código, guarda con la "
                "codificación equivocada y NO deja copia de respaldo; así se dañaron archivos antes. Usa escribir_archivo para "
                "crear o reemplazar un archivo completo (por partes con modo='agregar' si es largo) y editar_archivo para cambiar "
                "partes. PowerShell es para EJECUTAR, instalar y probar.")
    pre = _pre_backup(ctx, comando)
    cmd = ps_command(comando)
    if a.get("segundo_plano"):
        return run_background(cmd, ctx, comando[:200])
    out = run_streaming(cmd, ctx, timeout=min(int(a.get("timeout_segundos") or 300), 3600))
    cambiados = _post_backup(ctx, pre)
    if cambiados:
        out += f"\n(El comando modificó {', '.join(cambiados)}; quedó copia de respaldo: deshacer_cambio puede volver atrás.)"
    return out


def _snapshot(root: Path):
    try:
        return {p: p.stat().st_mtime for p in _iter_files(root)}
    except Exception:
        return {}


def t_ejecutar_python(a, ctx):
    pre = _pre_backup(ctx, a.get("codigo", ""))
    tmp = ctx.cwd / f".noviq_tmp_{uuid.uuid4().hex[:8]}.py"
    tmp.write_text(a.get("codigo", ""), encoding="utf-8")
    before = _snapshot(ctx.cwd)
    try:
        from .pruebas import project_python
        res = run_streaming([project_python(ctx.cwd), str(tmp)], ctx, timeout=min(int(a.get("timeout_segundos") or 600), 3600))
    finally:
        tmp.unlink(missing_ok=True)
    _post_backup(ctx, pre)
    after = _snapshot(ctx.cwd)
    new = [p for p, m in after.items() if (p not in before or before[p] != m) and ".noviq_tmp_" not in p.name]
    for p in sorted(new)[:12]:
        if p.suffix.lower() in IMG_EXT:
            ctx.emit("image", {"name": p.name, "path": str(p)})
        else:
            file_event(ctx, p, action="creado")
    if new:
        res += "\nArchivos creados/modificados: " + ", ".join(str(p.relative_to(ctx.cwd)) for p in new[:20])
    return res


def t_procesos(a, ctx):
    acc = a.get("accion", "listar")
    if acc == "listar":
        if not PROCS:
            return "No hay procesos en segundo plano."
        return "\n".join(f"id={k}  {'activo' if v['proc'].poll() is None else 'terminado(' + str(v['proc'].returncode) + ')'}  "
                         f"{int(time.time() - v['start'])}s  {v['cmd']}" for k, v in PROCS.items())
    pr = PROCS.get(a.get("id", ""))
    if not pr:
        return "ERROR: id de proceso desconocido. Usa accion='listar'."
    if acc == "salida":
        txt = pr["log"].read_bytes().decode("utf-8", errors="replace")
        st = "activo" if pr["proc"].poll() is None else f"terminado (código {pr['proc'].returncode})"
        return f"Estado: {st}\n" + cut(txt[-15000:], 15000)
    if acc == "detener":
        _kill_tree(pr["proc"])
        return f"Proceso {a.get('id')} detenido."
    return "ERROR: acción inválida (listar | salida | detener)."


# ═════════════════════════════ INTERNET ═════════════════════════════
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/128 Safari/537.36", "Accept-Language": "es-ES,es;q=0.9,en;q=0.8"}


class _Text(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "head", "nav", "footer", "iframe", "form"}

    def __init__(self):
        super().__init__()
        self.parts, self.skip, self.links = [], 0, []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in ("p", "br", "div", "li", "tr", "section", "article", "pre"):
            self.parts.append("\n")
        elif tag in ("h1", "h2", "h3", "h4"):
            self.parts.append("\n\n## ")
        elif tag == "a":
            href = dict(attrs).get("href")
            if href and href.startswith("http") and len(self.links) < 40:
                self.links.append(href)

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def t_leer_pagina_web(a, ctx):
    url = a.get("url", "")
    if not url.startswith("http"):
        url = "https://" + url
    r = requests.get(url, headers=UA, timeout=30)
    ctype = r.headers.get("content-type", "")
    if "html" not in ctype:
        return cut(f"[{r.status_code}] {ctype}\n" + r.text, 40000)
    ex = _Text()
    ex.feed(r.text)
    text = re.sub(r"\n\s*\n+", "\n\n", "".join(ex.parts))
    text = re.sub(r"[ \t]+", " ", text).strip()
    return cut(f"[{r.status_code}] {r.url}\n\n{text}", 40000)


def t_buscar_en_internet(a, ctx):
    q = a.get("consulta", "")
    r = requests.post("https://html.duckduckgo.com/html/", data={"q": q, "kl": "wt-wt"}, headers=UA, timeout=30)
    blocks = re.findall(r'<a[^>]+class="result__a"[^>]+href="([^"]+)"[^>]*>(.*?)</a>(.*?)(?=<a[^>]+class="result__a"|$)',
                        r.text, re.S)
    clean = lambda s: html.unescape(re.sub("<.*?>", "", s or "")).strip()
    out = []
    for href, title, rest in blocks[:10]:
        if "uddg=" in href:
            href = unquote(parse_qs(urlparse(href).query).get("uddg", [href])[0])
        snip = re.search(r'class="result__snippet"[^>]*>(.*?)</a>', rest, re.S)
        out.append(f"- {clean(title)}\n  {href}\n  {clean(snip.group(1) if snip else '')}")
    if out:
        return "\n".join(out)
    # respaldo: Bing
    r = requests.get("https://www.bing.com/search", params={"q": q}, headers=UA, timeout=30)
    for m in re.finditer(r'<li class="b_algo".*?<h2><a href="([^"]+)"[^>]*>(.*?)</a></h2>.*?(?:<p[^>]*>(.*?)</p>)?', r.text, re.S):
        out.append(f"- {clean(m.group(2))}\n  {m.group(1)}\n  {clean(m.group(3))}")
        if len(out) >= 10:
            break
    return "\n".join(out) if out else "Sin resultados. Prueba leer_pagina_web con una URL directa."


def t_descargar_archivo(a, ctx):
    url = a.get("url", "")
    name = a.get("ruta") or (Path(urlparse(url).path).name or "descarga.bin")
    p = ctx.resolve(name)
    p.parent.mkdir(parents=True, exist_ok=True)
    with requests.get(url, headers=UA, stream=True, timeout=60) as r:
        r.raise_for_status()
        with open(p, "wb") as f:
            for chunk in r.iter_content(1 << 16):
                f.write(chunk)
                if ctx.stop_ev.is_set():
                    return "Descarga cancelada."
    file_event(ctx, p, action="descargado")
    return f"Descargado: {p} ({p.stat().st_size:,} bytes)"


# ═════════════════════════════ IMÁGENES / PANTALLA ═════════════════════════════
SIZES = {"cuadrado": (1024, 1024), "horizontal": (1344, 768), "vertical": (768, 1344)}


def t_generar_imagen(a, ctx):
    model = CONFIG.get("image_model") or "black-forest-labs/flux.1-schnell"
    from . import vault
    keys = vault.get_keys("nvidia")
    if not keys:
        return "ERROR: generar imágenes necesita una API key de NVIDIA (⚙ Configuración)."
    w, h = SIZES.get(a.get("formato") or "cuadrado", (1024, 1024))
    prompt = a.get("prompt", "")
    if "schnell" in model:
        payload = {"prompt": prompt, "width": w, "height": h, "seed": 0, "steps": 4}
    elif "flux" in model:
        payload = {"prompt": prompt, "mode": "base", "cfg_scale": 3.5, "width": w, "height": h, "seed": 0, "steps": 30}
    else:
        payload = {"prompt": prompt, "width": w, "height": h}
    base = os.environ.get("NVIDIA_IMAGE_BASE", "https://ai.api.nvidia.com/v1/genai")
    url = f"{base}/{model}"
    headers = {"Authorization": f"Bearer {keys[0]}", "Accept": "application/json"}
    r = requests.post(url, headers=headers, json=payload, timeout=240)
    if r.status_code in (400, 422):
        r = requests.post(url, headers=headers, json={"prompt": prompt}, timeout=240)
    if r.status_code != 200:
        return f"ERROR {r.status_code} generando imagen: {r.text[:800]}"
    d = r.json()
    b64, reason = None, ""
    if isinstance(d.get("artifacts"), list) and d["artifacts"]:
        b64 = d["artifacts"][0].get("base64")
        reason = d["artifacts"][0].get("finishReason", "")
    b64 = b64 or d.get("image") or ((d.get("data") or [{}])[0].get("b64_json"))
    if reason == "CONTENT_FILTERED":
        return "ERROR: el filtro de NVIDIA bloqueó el prompt (a veces por palabras como 'avatar', 'brick', 'grand'). Reescríbelo."
    if not b64:
        return "ERROR: respuesta sin imagen: " + json.dumps(d)[:400]
    raw = base64.b64decode(b64)
    ext = ".png" if raw[:4] == b"\x89PNG" else ".jpg"
    name = re.sub(r"[^\w\-]+", "_", a.get("nombre_archivo") or "imagen")[:50]
    folder = ctx.cwd / "imagenes"
    folder.mkdir(exist_ok=True)
    p = folder / f"{name}_{datetime.now():%Y%m%d_%H%M%S}{ext}"
    p.write_bytes(raw)
    ctx.emit("image", {"name": p.name, "path": str(p)})
    return f"Imagen generada, mostrada al usuario y guardada en: {p}"


def t_captura_pantalla(a, ctx):
    try:
        from PIL import ImageGrab
    except ImportError:
        return "ERROR: falta Pillow (pip install pillow)."
    im = ImageGrab.grab(all_screens=bool(a.get("todas_las_pantallas")))
    folder = ctx.cwd / "capturas"
    folder.mkdir(exist_ok=True)
    p = folder / f"captura_{datetime.now():%Y%m%d_%H%M%S}.png"
    im.save(p)
    ctx.emit("image", {"name": p.name, "path": str(p)})
    return {"text": f"Captura guardada en {p} ({im.width}x{im.height}). Te la adjunto para que la veas.",
            "images": [img_to_dataurl(p)]}


def t_abrir(a, ctx):
    obj = (a.get("ruta_o_url") or "").strip()
    if re.match(r"https?://", obj):
        webbrowser.open(obj)
        return f"Abierto en el navegador: {obj}"
    p = ctx.resolve(obj)
    if not p.exists():
        return f"ERROR: no existe {p}"
    if IS_WINDOWS:
        os.startfile(str(p))  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(p)])
    else:
        subprocess.Popen(["xdg-open", str(p)])
    return f"Abierto con la aplicación predeterminada: {p}"


# ═════════════════════════════ TAREAS Y MEMORIA ═════════════════════════════
def t_actualizar_tareas(a, ctx):
    tareas = a.get("tareas") or []
    norm = []
    for t in tareas:
        if isinstance(t, str):
            t = {"contenido": t}
        norm.append({"contenido": str(t.get("contenido") or t.get("texto") or ""),
                     "estado": t.get("estado") if t.get("estado") in ("pendiente", "en_progreso", "completada") else "pendiente"})
    ctx.conv["tasks"] = norm
    ctx.emit("tasks", {"tasks": norm})
    done = sum(t["estado"] == "completada" for t in norm)
    return f"Lista de tareas actualizada ({done}/{len(norm)} completadas)."


def read_memory():
    return MEMORY_FILE.read_text(encoding="utf-8") if MEMORY_FILE.exists() else ""


def project_meta_dir(project):
    return meta_dir(project)


def read_project_memory(project):
    if not project:
        return ""
    f = meta_dir(project, create=False) / "memoria.md"
    return f.read_text(encoding="utf-8") if f.exists() else ""


def append_log(conv, line):
    """Bitácora: registro breve de lo hecho (conversación + proyecto) para que otra IA pueda continuar."""
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    conv.setdefault("bitacora", []).append(f"{stamp} {line}")
    conv["bitacora"] = conv["bitacora"][-300:]
    if conv.get("project") and project_path(conv["project"]).exists():
        try:
            with open(project_meta_dir(conv["project"]) / "bitacora.md", "a", encoding="utf-8") as f:
                f.write(f"- {stamp} [{conv.get('title', '')[:40]}] {line}\n")
        except OSError:
            pass


def read_project_log(project, n=25):
    if not project:
        return ""
    f = meta_dir(project, create=False) / "bitacora.md"
    if not f.exists():
        return ""
    return "\n".join(f.read_text(encoding="utf-8").splitlines()[-n:])


def t_memoria(a, ctx):
    acc = a.get("accion", "ver")
    proj = ctx.conv.get("project") if a.get("alcance") == "proyecto" else None
    if a.get("alcance") == "proyecto" and not proj:
        return "ERROR: no hay proyecto activo; usa alcance='usuario' o crea/abre un proyecto."
    f = (project_meta_dir(proj) / "memoria.md") if proj else MEMORY_FILE
    mem = f.read_text(encoding="utf-8") if f.exists() else ""
    donde = f"del proyecto {proj}" if proj else "del usuario"
    if acc == "ver":
        return mem or f"(memoria {donde} vacía)"
    if acc == "agregar":
        texto = (a.get("texto") or "").strip()
        f.write_text((mem.rstrip() + "\n" if mem.strip() else "") + f"- {texto}\n", encoding="utf-8")
        return f"Guardado en la memoria {donde}."
    if acc == "reemplazar_todo":
        f.write_text(a.get("texto") or "", encoding="utf-8")
        return f"Memoria {donde} reescrita."
    return "ERROR: acción inválida (ver | agregar | reemplazar_todo)."


# ═════════════════════════════ AYUDA DEL USUARIO Y NAVEGADOR ═════════════════════════════
ASK = {}  # id -> {"event", "texto", "imagenes"}


def t_pedir_ayuda_usuario(a, ctx):
    """Pide al usuario que haga algo manualmente (o dé un dato) y espera su respuesta y capturas."""
    ev = threading.Event()
    ASK[ctx.call_id] = {"event": ev, "texto": "", "imagenes": []}
    ctx.emit("ask_user", {"id": ctx.call_id, "instrucciones": a.get("instrucciones", ""),
                          "pide_captura": bool(a.get("pedir_captura", True))})
    t0 = time.time()
    while not ev.wait(0.5):
        if ctx.stop_ev.is_set() or time.time() - t0 > 3600:
            ASK.pop(ctx.call_id, None)
            return "El usuario no respondió (cancelado o tiempo agotado)."
    ans = ASK.pop(ctx.call_id, {})
    imgs = ans.get("imagenes") or []
    txt = (ans.get("texto") or "").strip() or "(sin texto)"
    return {"text": f"Respuesta del usuario: {txt}" + (f"\n[El usuario adjuntó {len(imgs)} captura(s)]" if imgs else ""),
            "images": imgs[:4]}


def t_agente_navegador(a, ctx):
    from . import subagente
    tarea = a.get("tarea", "")
    tools = subagente.browser_tools()
    if not tools:
        return ("ERROR: el control del navegador no está activo. El usuario puede activarlo en 🔌 MCP → «Chrome» o «Edge» "
                "(requiere Node.js). Mientras tanto usa pedir_ayuda_usuario para pedirle que haga los pasos en su "
                "navegador y te envíe capturas.")
    with subagente.recurso("navegador", ctx):
        res = subagente.run(ctx.conv, ctx, tarea, tools, subagente.BROWSER_SYSTEM, label="Agente navegador")
    if res is None:
        return ("ERROR: ningún modelo disponible pudo controlar el navegador (se necesitan herramientas nativas). "
                "Usa pedir_ayuda_usuario para que el usuario lo haga y te envíe capturas.")
    return "Resultado del agente navegador:\n" + res


def t_control_pc(a, ctx):
    from . import pc
    return pc.control_pc(a, ctx)


def t_agente_pc(a, ctx):
    """Delegado: un subagente maneja el escritorio (Windows-MCP si está conectado; si no, el control integrado)."""
    from . import pc, subagente
    from .mcp import MANAGER
    tarea = a.get("tarea", "")
    if MANAGER.has_pc():
        tools, executor, how = MANAGER.pc_tools(), None, "Windows-MCP"
    else:
        tools, how = [pc.TOOL_DEF], "control integrado"

        def executor(name, args):
            return pc.control_pc(args, ctx)
    with subagente.recurso("pc", ctx):
        res = subagente.run(ctx.conv, ctx, tarea, tools, pc.PC_SYSTEM, label=f"Agente del PC ({how})", executor=executor, max_steps=40)
    if res is None:
        return ("ERROR: ningún modelo disponible pudo manejar el PC (se necesita un modelo con herramientas nativas, ej. Gemini, "
                "GLM, Qwen3, Kimi). Puedes usar control_pc paso a paso tú mismo o pedir_ayuda_usuario.")
    return "Resultado del agente del PC:\n" + res


def t_programar_tarea(a, ctx):
    from . import programador
    return programador.tool(a, ctx)


def t_probar_proyecto(a, ctx):
    """Agente de pruebas: revisa sintaxis, nombres, dependencias y ejecuta las pruebas del proyecto."""
    from . import pruebas as PR
    dynamic = a.get("ejecutar_pruebas", True) is not False
    archivos = a.get("archivos") or None
    if isinstance(archivos, str):
        archivos = [archivos]
    res = PR.run_checks(ctx.cwd, only=archivos, dynamic=dynamic, allow_install=True, web=dynamic and not archivos,
                        emit=lambda t: ctx.emit("tool_output", {"id": ctx.call_id, "text": t + "\n"}))
    out = PR.report(res)
    # comando de pruebas propio del proyecto (npm test, dotnet test, mvn test…)
    meta = project_info(ctx.conv["project"]) if ctx.conv.get("project") else {}
    cmd = (meta.get("comandos") or {}).get("probar", "")
    if dynamic and res["ok"] and cmd and not cmd.startswith("python -m pytest"):
        r = run_streaming(ps_command(cmd), ctx, timeout=600)
        if not r.startswith("Código de salida: 0"):
            lines = [l for l in r.splitlines() if re.search(r"(?i)error|fail|exception|✖|×", l)][:8]
            loc = t_localizar_error({"texto_error": r}, ctx)
            out = (f"❌ `{cmd}` falló:\n" + "\n".join(lines)[:1200] +
                   ("\n\n" + loc[:2000] if not loc.startswith("No encontré") else ""))
        else:
            out += f"\n✅ `{cmd}` pasó."
    return out


def t_diagnosticar(a, ctx):
    """Diagnóstico completo del proyecto (o de algunos archivos): lista priorizada de problemas con cómo arreglarlos."""
    from . import diagnostico as D
    archivos = a.get("archivos") or None
    if isinstance(archivos, str):
        archivos = [archivos]
    root = ctx.cwd
    if a.get("ruta"):
        r = ctx.resolve(a["ruta"])
        if r.is_dir():
            root = r
        elif r.is_file():
            archivos = [str(r.relative_to(root)) if inside_root(r) and str(r).startswith(str(root)) else str(r)]
    pm = ctx.conv.get("permiso") or CONFIG.get("permission_mode", "preguntar")
    res = D.run(root, ctx=ctx, only=archivos, dinamico=a.get("abrir_navegador", True) is not False,
                ejecutar=bool(a.get("ejecutar")) and pm != "preguntar", reparar=a.get("reparar", True) is not False,
                emit=lambda t: ctx.emit("tool_output", {"id": ctx.call_id, "text": t + "\n"}))
    rep = D.reporte(res, max_items=int(a.get("max_problemas") or 10))
    w = res.get("web") or {}
    if a.get("captura") and w.get("png"):
        return {"text": rep, "images": ["data:image/png;base64," + base64.b64encode(w["png"]).decode()]}
    return rep


def t_probar_web(a, ctx):
    """Abre una página (archivo HTML del proyecto o URL de un servidor) en un navegador invisible y la prueba."""
    from . import navegador as NAV
    if not NAV.find_browser():
        return ("ERROR: no encontré Microsoft Edge, Google Chrome ni Chromium en este equipo para probar la página. "
                "Usa abrir para mostrarla al usuario y pedir_ayuda_usuario para que te diga qué ve.")
    acciones = a.get("acciones") or []
    if isinstance(acciones, str):
        try:
            acciones = json.loads(acciones)
        except Exception:
            acciones = [{"tipo": "tecla", "tecla": k.strip()} for k in acciones.split(",") if k.strip()]
    if isinstance(acciones, dict):
        acciones = [acciones]
    objetivo = a.get("ruta_o_url") or a.get("url") or a.get("ruta") or "index.html"
    ctx.emit("tool_output", {"id": ctx.call_id, "text": f"🌐 Abriendo {objetivo} en un navegador invisible…\n"})
    try:
        r = NAV.probar(objetivo, cwd=ctx.cwd, acciones=acciones, evaluar=a.get("evaluar_js") or "",
                       captura=True, espera_ms=int(a.get("espera_ms") or 2500))
    except Exception as e:
        return f"ERROR: no pude probar la página: {type(e).__name__}: {e}"
    rep = NAV.reporte(r)
    if r.get("png") and a.get("ver_captura", True) is not False:
        folder = ctx.cwd / "capturas"
        try:
            folder.mkdir(exist_ok=True)
            f = folder / f"prueba_web_{datetime.now():%Y%m%d_%H%M%S}.png"
            f.write_bytes(r["png"])
            for old in sorted(folder.glob("prueba_web_*.png"))[:-8]:  # solo se guardan las últimas capturas de prueba
                old.unlink(missing_ok=True)
            ctx.emit("image", {"name": f.name, "path": str(f)})
            rep += f"\n\n📸 Captura guardada en {f.relative_to(ctx.cwd)} (te la adjunto para que veas cómo se ve)."
            return {"text": rep, "images": [img_to_dataurl(f)]}
        except OSError:
            pass
    return rep


def t_planificar_en_equipo(a, ctx):
    from . import subagente
    from . import estado as E
    objetivo = a.get("objetivo", "")
    contexto = (a.get("contexto") or "") + "\n\n" + (E.build(ctx.conv) or "")
    props = subagente.plan_team(ctx.conv, ctx, objetivo, contexto)
    if not props:
        return "ERROR: ninguna IA del equipo respondió. Haz tú el plan con actualizar_tareas."
    txt = "\n\n".join(f"### Propuesta del {r} ({m.split('::')[-1]})\n{p}" for r, m, p in props)
    return (f"{len(props)} IAs propusieron planes. COMBINA lo mejor de cada una en UN plan final y regístralo con "
            "actualizar_tareas. Si hay partes independientes (archivos distintos), puedes repartirlas con trabajar_en_equipo.\n\n" + txt)


def t_trabajar_en_equipo(a, ctx):
    from . import subagente
    tareas = a.get("tareas") or []
    if isinstance(tareas, str):
        try:
            tareas = json.loads(tareas)
        except Exception:
            tareas = []
    return subagente.work_team(ctx.conv, ctx, tareas)


# ═════════════════════════════ DEFINICIONES ═════════════════════════════
def T(name, desc, props, required=()):
    return {"type": "function", "function": {"name": name, "description": desc, "parameters": {
        "type": "object", "properties": props, "required": list(required)}}}


S = lambda d: {"type": "string", "description": d}
B = lambda d: {"type": "boolean", "description": d}
I = lambda d: {"type": "integer", "description": d}

BUILTIN = [
    T("crear_proyecto", "Crea un proyecto nuevo (carpeta propia en Documentos\\Noviq_Proyectos, creada con PowerShell, con subcarpetas según el tipo, README, .gitignore y git) y lo activa como directorio de trabajo. Úsalo SIEMPRE antes de empezar un programa/app/web nuevo.",
      {"nombre": S("Nombre corto del proyecto, ej. 'Calculadora_Ventas'"),
       "tipo": {"type": "string", "enum": list(TIPOS), "description": "Tipo de proyecto"},
       "descripcion": S("Qué es el proyecto, una o dos frases")}, ["nombre", "tipo"]),
    T("cambiar_proyecto", "Activa un proyecto existente como directorio de trabajo (para seguir trabajando en él).",
      {"nombre": S("Nombre del proyecto")}, ["nombre"]),
    T("escribir_archivo", "Crea un archivo NUEVO con su contenido COMPLETO (código, HTML, docs, .bat, .ps1, etc.). Rutas relativas = proyecto activo. Para cambiar o corregir archivos existentes usa editar_archivo (no reescribas todo).",
      {"ruta": S("Ruta, ej. 'src/main.py'"), "contenido": S("Contenido del archivo"),
       "modo": {"type": "string", "enum": ["crear", "agregar"], "description": "'crear' (por defecto) o 'agregar' al final (para escribir archivos largos por partes)"},
       "sobrescribir": B("true solo si de verdad quieres reemplazar por completo un archivo existente grande")},
      ["ruta", "contenido"]),
    T("editar_archivo", "Modifica SOLO las partes necesarias de un archivo existente (edición quirúrgica, con copia de respaldo). Opciones: (a) texto_anterior + texto_nuevo (copia texto_anterior del archivo SIN los números de línea; tolera diferencias de espacios y sangría); (b) 'ediciones': lista de {texto_anterior, texto_nuevo} para varios cambios a la vez; (c) desde_linea/hasta_linea + texto_nuevo para REEMPLAZAR ese rango de líneas (lo más seguro si ya leíste las líneas); (d) despues_de_linea + texto_nuevo para INSERTAR código nuevo (0 = al principio). Si falla, el error te muestra las líneas reales y sus números.",
      {"ruta": S("Archivo a editar"), "texto_anterior": S("Texto que existe en el archivo (sin números de línea)"),
       "texto_nuevo": S("Texto que lo reemplaza (o que se inserta)"), "reemplazar_todo": B("Reemplazar todas las apariciones"),
       "ediciones": {"type": "array", "description": "Varias ediciones en orden", "items": {"type": "object", "properties": {
           "texto_anterior": {"type": "string"}, "texto_nuevo": {"type": "string"}}, "required": ["texto_anterior", "texto_nuevo"]}},
       "desde_linea": I("Primera línea a reemplazar (modo rango)"), "hasta_linea": I("Última línea a reemplazar (modo rango)"),
       "despues_de_linea": I("Insertar texto_nuevo después de esta línea (0 = al principio)")},
      ["ruta"]),
    T("diagnosticar", "🩺 Diagnóstico COMPLETO del proyecto como lo haría un programador experto: codificación dañada (la repara), código incompleto ('...'), funciones que se llaman pero no existen o no se cargan, orden de los <script>, ids de HTML que no existen, librerías (THREE, jQuery…) sin cargar, sintaxis, nombres no definidos, dependencias, y ABRE la página en un navegador invisible para ver los errores reales de la consola y si la pantalla queda vacía. Devuelve la lista PRIORIZADA de problemas con archivo:línea y cómo arreglar cada uno. Úsalo PRIMERO cuando algo no funciona o el usuario pide revisar/arreglar.",
      {"ruta": S("Carpeta o archivo (vacío = proyecto activo)"),
       "archivos": {"type": "array", "items": {"type": "string"}, "description": "Solo estos archivos (opcional)"},
       "abrir_navegador": B("Abrir la página web en el navegador de pruebas (por defecto sí)"),
       "ejecutar": B("Ejecutar también el programa (Python/Node) para ver si arranca"),
       "captura": B("Adjuntar la captura de la página")}),
    T("probar_web", "🌐 Prueba una página web, juego o app web como un usuario real, en un navegador invisible: errores de JavaScript con archivo:línea, archivos que no cargan (404/CDN), mensajes de consola, captura de pantalla (la verás) y si queda vacía. Puede simular acciones: teclas (flechas, espacio, WASD), clics, arrastrar el mouse, rueda, escribir texto y ejecutar JavaScript para leer el estado (ej. puntaje). Úsalo para VERIFICAR que una web o juego funciona antes de decir que está listo.",
      {"ruta_o_url": S("Archivo HTML del proyecto (ej. 'index.html') o URL (ej. http://localhost:8000)"),
       "acciones": {"type": "array", "description": "Pasos en orden. Ej: [{\"tipo\":\"clic\",\"selector\":\"#jugar\"},{\"tipo\":\"tecla\",\"tecla\":\"ArrowLeft\",\"veces\":3},{\"tipo\":\"arrastrar\",\"x\":600,\"y\":400,\"x2\":800,\"y2\":400},{\"tipo\":\"esperar\",\"ms\":1000},{\"tipo\":\"js\",\"codigo\":\"score\"}]",
                    "items": {"type": "object", "properties": {
                        "tipo": {"type": "string", "enum": ["tecla", "clic", "doble_clic", "arrastrar", "rueda", "escribir", "esperar", "js"]},
                        "tecla": {"type": "string"}, "veces": {"type": "integer"}, "selector": {"type": "string"},
                        "x": {"type": "number"}, "y": {"type": "number"}, "x2": {"type": "number"}, "y2": {"type": "number"},
                        "texto": {"type": "string"}, "ms": {"type": "integer"}, "codigo": {"type": "string"}, "delta": {"type": "number"}},
                        "required": ["tipo"]}},
       "evaluar_js": S("Expresión JavaScript a evaluar al final (ej. 'score' o 'document.title')"),
       "espera_ms": I("Milisegundos a esperar tras cargar (def. 2500)")}),
    T("localizar_error", "Analiza un mensaje de error o traza (Python, JS, C#, Java, PowerShell, PHP…) y devuelve los archivos y líneas exactas del proyecto donde ocurre, con el código alrededor. Úsalo PRIMERO para corregir errores.",
      {"texto_error": S("Mensaje de error o traza completa")}, ["texto_error"]),
    T("mapa_codigo", "Muestra el mapa del proyecto: cada archivo con sus funciones/clases y número de línea. Útil para saber qué archivo y qué parte modificar sin leerlo todo.",
      {"ruta": S("Carpeta o archivo (vacío = proyecto)")}),
    T("deshacer_cambio", "Restaura un archivo a su versión anterior (antes del último cambio hecho por la IA).",
      {"ruta": S("Archivo a restaurar")}, ["ruta"]),
    T("leer_archivo", "Lee un archivo: texto/código (con números de línea), PDF, Word (.docx), Excel (.xlsx) o imagen (la verás).",
      {"ruta": S("Ruta del archivo"), "desde_linea": I("Línea inicial (opcional)"), "hasta_linea": I("Línea final (opcional)")},
      ["ruta"]),
    T("listar_directorio", "Lista archivos y carpetas. Vacío = proyecto activo.",
      {"ruta": S("Carpeta"), "recursivo": B("Incluir subcarpetas (árbol)")}),
    T("buscar_archivos", "Busca archivos por nombre con comodines (glob), ej. '*.py' o 'src/*.js'.",
      {"patron": S("Patrón glob"), "ruta": S("Carpeta raíz (opcional)")}, ["patron"]),
    T("buscar_en_archivos", "Busca texto o expresión regular dentro de los archivos (como grep). Devuelve archivo:línea: texto.",
      {"regex": S("Texto o regex"), "ruta": S("Carpeta o archivo (opcional)"), "incluir": S("Filtro glob, ej. '*.py'"),
       "mayusculas": B("Distinguir mayúsculas"), "contexto": I("Líneas de contexto antes/después (0-6)")}, ["regex"]),
    T("mover_o_copiar", "Mueve/renombra o copia archivos o carpetas.",
      {"origen": S("Ruta origen"), "destino": S("Ruta destino"), "copiar": B("true = copiar, false = mover")},
      ["origen", "destino"]),
    T("eliminar", "Envía un archivo o carpeta a la Papelera de reciclaje.", {"ruta": S("Ruta a eliminar")}, ["ruta"]),
    T("ejecutar_powershell", f"Ejecuta PowerShell en la PC del usuario ({OS_NAME}), en la carpeta del proyecto activo: ejecutar programas, instalar, compilar, probar, git. La salida se ve en vivo. Para servidores o procesos largos usa segundo_plano=true. NO lo uses para escribir archivos de código (usa escribir_archivo/editar_archivo).",
      {"comando": S("Comando o script de PowerShell"), "explicacion": S("Qué hace, en una frase, para el usuario"),
       "segundo_plano": B("Ejecutar sin esperar (servidores, watchers)"), "timeout_segundos": I("Máximo de espera (def. 300)")},
      ["comando", "explicacion"]),
    T("ejecutar_python", "Ejecuta código Python (cálculos, gráficos, Excel/Word/PDF, datos, automatizaciones) en la carpeta del proyecto. Puede instalar paquetes con subprocess + pip. Archivos e imágenes nuevos se muestran al usuario.",
      {"codigo": S("Código Python completo"), "explicacion": S("Qué hace, en una frase"), "timeout_segundos": I("Máximo de espera")},
      ["codigo", "explicacion"]),
    T("procesos", "Gestiona procesos en segundo plano: listar, ver salida o detener.",
      {"accion": {"type": "string", "enum": ["listar", "salida", "detener"]}, "id": S("Id del proceso")}, ["accion"]),
    T("buscar_en_internet", "Busca en internet y devuelve títulos, enlaces y resúmenes.", {"consulta": S("Consulta")}, ["consulta"]),
    T("leer_pagina_web", "Descarga una página web y devuelve su texto.", {"url": S("URL")}, ["url"]),
    T("descargar_archivo", "Descarga un archivo de internet al proyecto.", {"url": S("URL"), "ruta": S("Dónde guardarlo (opcional)")}, ["url"]),
    T("generar_imagen", "Dibuja/genera una imagen con IA (FLUX de NVIDIA) y la muestra. El prompt DEBE estar en inglés y ser detallado.",
      {"prompt": S("Descripción detallada en INGLÉS"), "formato": {"type": "string", "enum": list(SIZES)},
       "nombre_archivo": S("Nombre corto sin extensión")}, ["prompt"]),
    T("captura_pantalla", "Toma una captura de la pantalla del usuario y te la muestra (para ver qué hay en pantalla o revisar una app/web abierta).",
      {"todas_las_pantallas": B("Capturar todos los monitores")}),
    T("abrir", "Abre un archivo con su programa predeterminado o una URL en el navegador (ej. para mostrar una página web creada).",
      {"ruta_o_url": S("Ruta o URL")}, ["ruta_o_url"]),
    T("actualizar_tareas", "Crea/actualiza tu lista de tareas visible para el usuario. Úsala en trabajos de varios pasos: planifica, marca 'en_progreso' la actual y 'completada' al terminar.",
      {"tareas": {"type": "array", "items": {"type": "object", "properties": {
          "contenido": {"type": "string"}, "estado": {"type": "string", "enum": ["pendiente", "en_progreso", "completada"]}},
          "required": ["contenido", "estado"]}}}, ["tareas"]),
    T("probar_proyecto", "Agente de pruebas: revisa sintaxis, nombres no definidos y dependencias (las instala), ejecuta las pruebas del proyecto y, si es una web, la abre en un navegador invisible. Devuelve un reporte corto con archivo:línea, el código alrededor y cómo arreglarlo. Úsalo después de escribir o corregir código.",
      {"archivos": {"type": "array", "items": {"type": "string"}, "description": "Solo estos archivos (vacío = todo el proyecto)"},
       "ejecutar_pruebas": B("Ejecutar también las pruebas (pytest, npm test…); por defecto sí")}),
    T("planificar_en_equipo", "Pide a VARIAS IAs (arquitecto, experto en calidad, desarrollador pragmático) que propongan un plan en paralelo; luego tú combinas lo mejor en un plan final. Úsalo en proyectos medianos o grandes, o cuando el problema es difícil.",
      {"objetivo": S("Qué hay que lograr, detallado"), "contexto": S("Requisitos, restricciones o datos útiles")}, ["objetivo"]),
    T("trabajar_en_equipo", "Reparte tareas INDEPENDIENTES del plan entre varios agentes de IA que trabajan EN PARALELO (cada uno con su propio modelo). Asigna a cada tarea sus archivos para que no se pisen. Tú, como coordinador, luego integras, pruebas y corriges.",
      {"tareas": {"type": "array", "items": {"type": "object", "properties": {
          "titulo": {"type": "string"}, "rol": {"type": "string", "description": "ej. 'Backend', 'Interfaz', 'Pruebas'"},
          "instrucciones": {"type": "string", "description": "Qué debe hacer exactamente, con nombres de funciones/rutas acordados"},
          "archivos": {"type": "array", "items": {"type": "string"}, "description": "Archivos que SOLO este agente puede crear/modificar"},
          "contexto": {"type": "string", "description": "Partes del plan que necesita conocer (interfaces entre archivos)"}},
          "required": ["titulo", "instrucciones", "archivos"]}}}, ["tareas"]),
    T("agente_navegador", "Delega una tarea en el navegador web (Chrome/Edge) a un agente especializado que abre páginas, hace clic, escribe, llena formularios y lee datos; devuelve un resumen. Describe la tarea completa y concreta (URL, qué hacer, qué datos traer).",
      {"tarea": S("Tarea detallada para el navegador")}, ["tarea"]),
    T("agente_pc", "Delega una tarea en el ESCRITORIO de Windows a un agente especializado: abrir programas (Excel, Word, Bloc de notas, apps instaladas), hacer clic en botones y menús, escribir, usar atajos y verificar el resultado; devuelve un resumen. Describe la tarea completa y concreta.",
      {"tarea": S("Tarea detallada para hacer en el PC")}, ["tarea"]),
    T("control_pc", "Control directo y paso a paso del escritorio de Windows (mismo método que el navegador): accion='ventanas' o 'elementos' para leer (cada control tiene ref eN), luego 'clic'/'escribir' con ref, 'abrir' programas, 'enfocar' ventanas, 'teclas' (ctrl+s…), 'captura'. Para tareas largas usa agente_pc.",
      {"accion": {"type": "string", "enum": ["ventanas", "elementos", "abrir", "enfocar", "clic", "doble_clic", "clic_derecho",
                                              "escribir", "teclas", "esperar", "captura"]},
       "ref": S("Referencia eN de 'elementos'"), "x": I("Coordenada X"), "y": I("Coordenada Y"),
       "objetivo": S("Programa/ruta/URL para abrir, o título de la ventana para enfocar"), "texto": S("Texto a escribir"),
       "enter": B("Pulsar Enter tras escribir"), "teclas": S("Atajo: ctrl+s, alt+f4, enter…"), "segundos": {"type": "number"}}, ["accion"]),
    T("programar_tarea", "Programa un trabajo para que Noviark lo haga solo más tarde o repetidamente (una vez, cada día, ciertos días o cada N minutos), o lista/pausa/reanuda/elimina las tareas programadas.",
      {"accion": {"type": "string", "enum": ["crear", "listar", "eliminar", "pausar", "reanudar"]},
       "nombre": S("Nombre corto"), "instruccion": S("Qué debe hacer la IA en cada ejecución (completo)"),
       "tipo": {"type": "string", "enum": ["una_vez", "diario", "semanal", "intervalo"]},
       "fecha": S("una_vez: AAAA-MM-DDTHH:MM"), "hora": S("diario/semanal: HH:MM"),
       "dias": {"type": "array", "items": {"type": "integer"}, "description": "0=lunes … 6=domingo"},
       "cada_min": I("intervalo en minutos"), "id": S("id para eliminar/pausar/reanudar")}, ["accion"]),
    T("ejecutar_linux", "Ejecuta comandos bash en Linux (WSL de Windows) en la carpeta del proyecto: compilar con gcc/make, usar herramientas de Linux, scripts .sh, apt (con como_root=true). Si Linux no está instalado, lo indica.",
      {"comando": S("Comando bash"), "como_root": B("Ejecutar como root (para apt-get install, sin contraseña)"),
       "segundo_plano": B("Servidores/procesos largos"), "timeout_segundos": I("Máximo de espera")}, ["comando"]),
    T("contenedor", "Ejecuta un comando dentro de un contenedor Docker limpio con el proyecto montado en /work: prueba en un sistema aislado (otra versión de Python/Node, Linux puro) sin tocar la PC.",
      {"imagen": S("Imagen, ej. python:3.12-slim, node:22, ubuntu:24.04"), "comando": S("Comando sh"),
       "red": B("Permitir internet (por defecto sí)"), "puertos": {"type": "array", "items": {"type": "string"}, "description": "ej. ['8000:8000']"},
       "timeout_segundos": I("Máximo de espera")}, ["comando"]),
    T("entorno", "Entornos del proyecto: accion='info' (qué hay instalado: Python/.venv, Node, git, Docker, Linux), 'crear' (entorno virtual .venv o package.json), 'instalar' (paquetes o requirements.txt), 'congelar' (requirements.txt), 'instalar_linux' (instala WSL+Ubuntu).",
      {"accion": {"type": "string", "enum": ["info", "crear", "instalar", "congelar", "instalar_linux"]},
       "lenguaje": {"type": "string", "enum": ["python", "node"]}, "paquetes": {"type": "array", "items": {"type": "string"}}}, ["accion"]),
    T("instalar_herramienta", "Agente instalador: instala programas que falten (git, node, python, java, dotnet, go, rust, ffmpeg, docker, cmake, php, gh, uv…) con winget en Windows, apt en Linux/WSL, o paquetes pip/npm globales.",
      {"nombre": S("Nombre del programa o paquete"), "gestor": {"type": "string", "enum": ["auto", "winget", "apt", "pip", "npm"]}}, ["nombre"]),
    T("puntos_restauracion", "Puntos de restauración del proyecto (se crean solos antes de cada trabajo): listar, crear, ver diferencias o VOLVER a un punto si algo se rompió.",
      {"accion": {"type": "string", "enum": ["listar", "crear", "diferencias", "volver"]}, "id": S("Id del punto"),
       "mensaje": S("Descripción (crear)"), "archivos": {"type": "array", "items": {"type": "string"}, "description": "Solo estos archivos (opcional)"}}, ["accion"]),
    T("peticion_http", "Hace una petición HTTP (GET/POST/PUT/DELETE…) para probar APIs o el servidor que creaste (ej. http://127.0.0.1:5000/api). Devuelve estado, cabeceras y cuerpo.",
      {"url": S("URL"), "metodo": S("GET, POST…"), "cabeceras": {"type": "object"}, "cuerpo": {"description": "JSON u otro texto"},
       "timeout_segundos": I("Máximo de espera")}, ["url"]),
    T("pedir_ayuda_usuario", "Pide al usuario que haga algo que tú no puedes (iniciar sesión, un paso manual en el navegador o en un programa, conectar un equipo) o que te dé un dato; espera su respuesta y las capturas de pantalla que envíe.",
      {"instrucciones": S("Instrucciones claras y numeradas de lo que debe hacer el usuario"),
       "pedir_captura": B("Pedirle que adjunte una captura del resultado")}, ["instrucciones"]),
    T("memoria", "Memoria permanente. alcance='usuario': datos del usuario (preferencias, nombres, equipos). alcance='proyecto': decisiones, estado, pendientes y cómo ejecutar el proyecto activo (la leerá cualquier IA que continúe el trabajo).",
      {"accion": {"type": "string", "enum": ["ver", "agregar", "reemplazar_todo"]},
       "alcance": {"type": "string", "enum": ["usuario", "proyecto"]}, "texto": S("Texto a guardar")}, ["accion"]),
]

FUNCS = {
    "crear_proyecto": t_crear_proyecto, "cambiar_proyecto": t_cambiar_proyecto,
    "escribir_archivo": t_escribir_archivo, "editar_archivo": t_editar_archivo, "leer_archivo": t_leer_archivo,
    "listar_directorio": t_listar_directorio, "buscar_archivos": t_buscar_archivos,
    "buscar_en_archivos": t_buscar_en_archivos, "mover_o_copiar": t_mover_o_copiar, "eliminar": t_eliminar,
    "ejecutar_powershell": t_ejecutar_powershell, "ejecutar_python": t_ejecutar_python, "procesos": t_procesos,
    "buscar_en_internet": t_buscar_en_internet, "leer_pagina_web": t_leer_pagina_web,
    "descargar_archivo": t_descargar_archivo, "generar_imagen": t_generar_imagen,
    "captura_pantalla": t_captura_pantalla, "abrir": t_abrir, "actualizar_tareas": t_actualizar_tareas,
    "memoria": t_memoria, "deshacer_cambio": t_deshacer_cambio, "localizar_error": t_localizar_error,
    "mapa_codigo": t_mapa_codigo, "agente_navegador": t_agente_navegador,
    "planificar_en_equipo": t_planificar_en_equipo, "probar_proyecto": t_probar_proyecto, "trabajar_en_equipo": t_trabajar_en_equipo, "pedir_ayuda_usuario": t_pedir_ayuda_usuario,
    "agente_pc": t_agente_pc, "control_pc": t_control_pc, "programar_tarea": t_programar_tarea,
    "diagnosticar": t_diagnosticar, "probar_web": t_probar_web,
    "ejecutar_linux": lambda a, c: _ent().t_ejecutar_linux(a, c), "contenedor": lambda a, c: _ent().t_contenedor(a, c),
    "entorno": lambda a, c: _ent().t_entorno(a, c), "instalar_herramienta": lambda a, c: _ent().t_instalar_herramienta(a, c),
    "puntos_restauracion": lambda a, c: _ent().t_puntos_restauracion(a, c), "peticion_http": lambda a, c: _ent().t_peticion_http(a, c),
}


def _ent():
    from . import entornos
    return entornos

LABELS = {
    "crear_proyecto": "🗂 Creando proyecto", "cambiar_proyecto": "🗂 Abriendo proyecto",
    "escribir_archivo": "📄 Escribiendo archivo", "editar_archivo": "✏️ Editando archivo", "leer_archivo": "📖 Leyendo",
    "listar_directorio": "📁 Listando", "buscar_archivos": "🔍 Buscando archivos", "buscar_en_archivos": "🔍 Buscando texto",
    "mover_o_copiar": "📦 Moviendo/copiando", "eliminar": "🗑 Eliminando", "ejecutar_powershell": "⚡ PowerShell",
    "ejecutar_python": "🐍 Python", "procesos": "⚙ Procesos", "buscar_en_internet": "🔎 Buscando en internet",
    "leer_pagina_web": "🌐 Leyendo página", "descargar_archivo": "⬇ Descargando", "generar_imagen": "🎨 Generando imagen",
    "captura_pantalla": "🖥 Captura de pantalla", "abrir": "↗ Abriendo", "actualizar_tareas": "✅ Tareas",
    "memoria": "🧠 Memoria", "deshacer_cambio": "↩ Deshaciendo cambio", "localizar_error": "🎯 Localizando error",
    "mapa_codigo": "🗺 Mapa del código", "probar_proyecto": "🧪 Agente de pruebas", "planificar_en_equipo": "🧠 Plan en equipo (varias IAs)",
    "trabajar_en_equipo": "👥 Equipo de agentes trabajando", "agente_navegador": "🌐 Agente navegador", "pedir_ayuda_usuario": "🙋 Ayuda del usuario",
    "agente_pc": "🖥 Agente del PC", "control_pc": "🖱 Control del PC", "programar_tarea": "⏰ Tarea programada",
    "ejecutar_linux": "🐧 Linux", "contenedor": "🐳 Contenedor", "entorno": "🧰 Entorno", "instalar_herramienta": "📦 Instalador",
    "puntos_restauracion": "⏪ Puntos de restauración", "peticion_http": "🌐 Petición HTTP",
    "diagnosticar": "🩺 Diagnóstico del proyecto", "probar_web": "🌐 Probando la web en el navegador",
}

DANGER = {"ejecutar_powershell", "ejecutar_python", "eliminar", "ejecutar_linux", "contenedor", "instalar_herramienta"}
WRITE_PATH_ARGS = {"escribir_archivo": "ruta", "editar_archivo": "ruta", "mover_o_copiar": "destino", "deshacer_cambio": "ruta",
                   "descargar_archivo": "ruta"}


CMD_TOOLS = {"ejecutar_powershell", "ejecutar_python", "ejecutar_linux", "contenedor", "instalar_herramienta"}


def needs_permission(name, args, ctx):
    mode = ctx.conv.get("permiso") or CONFIG.get("permission_mode", "preguntar")
    if name == "ejecutar_powershell" and shell_file_write(args.get("comando", "")) and not args.get("forzar"):
        return False  # se bloquea igual (escribir código con PowerShell): no tiene sentido pedir permiso
    if name in CMD_TOOLS or name == "eliminar":
        # decisor «System One»: aun en modo automático, lo peligroso/irreversible siempre se pregunta
        try:
            from . import decisor
            nivel, fuente = decisor.riesgo(name, args)
        except Exception:
            nivel, fuente = "medio", "reglas"
        ctx.riesgo = (nivel, fuente)
        if nivel == "alto":
            return True
        if mode == "auto" or name in (ctx.conv.get("allowed_tools") or []):
            return False
        if mode == "auto_seguro":
            return nivel != "bajo" or name == "eliminar"
        return True
    if mode == "auto" or name in (ctx.conv.get("allowed_tools") or []):
        return False
    if name in DANGER:
        return True
    if name.startswith("mcp__"):
        return mode == "preguntar"
    if name in ("captura_pantalla", "agente_navegador", "probar_proyecto", "programar_tarea"):
        return mode == "preguntar"
    if name == "entorno":
        return args.get("accion") != "info"
    if name == "puntos_restauracion":
        return args.get("accion") == "volver" and mode == "preguntar"
    if name == "peticion_http":
        return (args.get("metodo") or "GET").upper() != "GET" and mode == "preguntar"
    if name in ("agente_pc", "control_pc"):
        # leer la pantalla no cambia nada; actuar sobre el escritorio sí pide permiso (salvo modo automático)
        return not (name == "control_pc" and args.get("accion") in ("ventanas", "elementos", "esperar"))
    if name in WRITE_PATH_ARGS:
        path = args.get(WRITE_PATH_ARGS[name]) or "."
        if name == "mover_o_copiar" and not inside_root(ctx.resolve(args.get("origen") or ".")):
            return True
        return not inside_root(ctx.resolve(path))
    return False
